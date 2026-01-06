"""
Excel Reader Tool
Tool for reading contact information from Excel files
"""

import os
from typing import Any, Dict, List, Optional
from datetime import datetime

from tools.base import Tool

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False
    try:
        import openpyxl
        OPENPYXL_AVAILABLE = True
    except ImportError:
        OPENPYXL_AVAILABLE = False


class ExcelReaderTool(Tool):
    """Tool that reads contact information from Excel files"""
    
    def __init__(self):
        super().__init__()
        self.name = "excel_reader"
        self.description = (
            "Reads contact information (name, email, phone) from Excel files. "
            "Supports .xlsx and .xls formats. Can also extract URLs for scraping."
        )
        
        if not PANDAS_AVAILABLE and not OPENPYXL_AVAILABLE:
            print("[ExcelReaderTool] Warning: pandas or openpyxl not installed. Install with: pip install pandas openpyxl")
    
    def _read_with_pandas(self, file_path: str, sheet_name: Optional[str] = None) -> List[Dict]:
        """Read Excel file using pandas"""
        try:
            if sheet_name:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
            else:
                df = pd.read_excel(file_path)
            
            contacts = []
            
            # Try to auto-detect columns
            name_col = None
            email_col = None
            phone_col = None
            url_col = None
            
            # Look for common column names
            for col in df.columns:
                col_lower = str(col).lower()
                if not name_col and any(x in col_lower for x in ['name', 'full name', 'contact name']):
                    name_col = col
                if not email_col and any(x in col_lower for x in ['email', 'e-mail', 'email address']):
                    email_col = col
                if not phone_col and any(x in col_lower for x in ['phone', 'telephone', 'mobile', 'contact']):
                    phone_col = col
                if not url_col and any(x in col_lower for x in ['url', 'website', 'link', 'source']):
                    url_col = col
            
            # If no specific columns found, use first few columns
            if not name_col and len(df.columns) > 0:
                name_col = df.columns[0]
            if not email_col and len(df.columns) > 1:
                email_col = df.columns[1]
            if not phone_col and len(df.columns) > 2:
                phone_col = df.columns[2]
            
            # Extract contacts
            for idx, row in df.iterrows():
                contact = {
                    'name': str(row[name_col]).strip() if name_col and pd.notna(row.get(name_col)) else None,
                    'email': str(row[email_col]).strip() if email_col and pd.notna(row.get(email_col)) else None,
                    'phone': str(row[phone_col]).strip() if phone_col and pd.notna(row.get(phone_col)) else None,
                    'source_url': str(row[url_col]).strip() if url_col and pd.notna(row.get(url_col)) else None,
                    'source_file': file_path,
                    'row_number': idx + 2  # +2 because Excel is 1-indexed and has header
                }
                
                # Clean up None values
                if contact['name'] == 'None' or contact['name'] == 'nan':
                    contact['name'] = None
                if contact['email'] == 'None' or contact['email'] == 'nan':
                    contact['email'] = None
                if contact['phone'] == 'None' or contact['phone'] == 'nan':
                    contact['phone'] = None
                if contact['source_url'] == 'None' or contact['source_url'] == 'nan':
                    contact['source_url'] = None
                
                # Only add if has at least one contact field
                if contact['name'] or contact['email'] or contact['phone'] or contact['source_url']:
                    contacts.append(contact)
            
            return contacts
            
        except Exception as e:
            print(f"[ExcelReaderTool] Error reading with pandas: {e}")
            return []
    
    def _read_with_openpyxl(self, file_path: str, sheet_name: Optional[str] = None) -> List[Dict]:
        """Read Excel file using openpyxl (fallback)"""
        try:
            from openpyxl import load_workbook
            
            wb = load_workbook(file_path, data_only=True)
            
            if sheet_name:
                ws = wb[sheet_name]
            else:
                ws = wb.active
            
            contacts = []
            
            # Read header row
            headers = [cell.value for cell in ws[1]]
            
            # Find column indices
            name_idx = None
            email_idx = None
            phone_idx = None
            url_idx = None
            
            for i, header in enumerate(headers):
                if header:
                    header_lower = str(header).lower()
                    if not name_idx and any(x in header_lower for x in ['name', 'full name', 'contact name']):
                        name_idx = i
                    if not email_idx and any(x in header_lower for x in ['email', 'e-mail', 'email address']):
                        email_idx = i
                    if not phone_idx and any(x in header_lower for x in ['phone', 'telephone', 'mobile', 'contact']):
                        phone_idx = i
                    if not url_idx and any(x in header_lower for x in ['url', 'website', 'link', 'source']):
                        url_idx = i
            
            # Default to first columns if not found
            if name_idx is None:
                name_idx = 0
            if email_idx is None and len(headers) > 1:
                email_idx = 1
            if phone_idx is None and len(headers) > 2:
                phone_idx = 2
            
            # Read data rows
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                contact = {
                    'name': str(row[name_idx]).strip() if name_idx < len(row) and row[name_idx] else None,
                    'email': str(row[email_idx]).strip() if email_idx and email_idx < len(row) and row[email_idx] else None,
                    'phone': str(row[phone_idx]).strip() if phone_idx and phone_idx < len(row) and row[phone_idx] else None,
                    'source_url': str(row[url_idx]).strip() if url_idx and url_idx < len(row) and row[url_idx] else None,
                    'source_file': file_path,
                    'row_number': row_idx
                }
                
                # Clean up
                if contact['name'] == 'None':
                    contact['name'] = None
                if contact['email'] == 'None':
                    contact['email'] = None
                if contact['phone'] == 'None':
                    contact['phone'] = None
                if contact['source_url'] == 'None':
                    contact['source_url'] = None
                
                # Only add if has data
                if contact['name'] or contact['email'] or contact['phone'] or contact['source_url']:
                    contacts.append(contact)
            
            return contacts
            
        except Exception as e:
            print(f"[ExcelReaderTool] Error reading with openpyxl: {e}")
            return []
    
    def run(
        self,
        file_path: str,
        sheet_name: Optional[str] = None,
        extract_urls_only: bool = False,
        **kwargs
    ) -> Dict[str, Any]:
        """
        Read contact information from Excel file
        
        Args:
            file_path: Path to Excel file (.xlsx or .xls)
            sheet_name: Specific sheet name to read (optional)
            extract_urls_only: If True, only extract URLs for scraping
        
        Returns:
            Dictionary with extracted contacts or URLs
        """
        if not os.path.exists(file_path):
            return {
                'success': False,
                'error': f'File not found: {file_path}',
                'contacts': [],
                'urls': []
            }
        
        if not PANDAS_AVAILABLE and not OPENPYXL_AVAILABLE:
            return {
                'success': False,
                'error': 'pandas or openpyxl not installed. Install with: pip install pandas openpyxl',
                'contacts': [],
                'urls': []
            }
        
        # Read file
        if PANDAS_AVAILABLE:
            contacts = self._read_with_pandas(file_path, sheet_name)
        else:
            contacts = self._read_with_openpyxl(file_path, sheet_name)
        
        # Extract URLs if requested
        urls = []
        if extract_urls_only:
            for contact in contacts:
                if contact.get('source_url'):
                    urls.append(contact['source_url'])
            return {
                'success': True,
                'tool': self.name,
                'file_path': file_path,
                'urls': list(set(urls)),  # Deduplicate
                'total_urls': len(set(urls)),
                'timestamp': datetime.now().isoformat()
            }
        
        # Return contacts
        # Deduplicate by email/phone
        unique_contacts = []
        seen = set()
        for contact in contacts:
            key = (
                contact.get('email', '').lower() if contact.get('email') else '',
                contact.get('phone', '') if contact.get('phone') else ''
            )
            if key not in seen and (contact.get('email') or contact.get('phone') or contact.get('name')):
                seen.add(key)
                unique_contacts.append(contact)
        
        return {
            'success': True,
            'tool': self.name,
            'file_path': file_path,
            'sheet_name': sheet_name,
            'total_contacts': len(unique_contacts),
            'contacts': unique_contacts,
            'timestamp': datetime.now().isoformat()
        }

